"""Pruebas de la exportación Excel basada en resultados ya resueltos."""

from __future__ import annotations

from contextlib import contextmanager
from copy import deepcopy
import importlib.util
from pathlib import Path

from openpyxl import load_workbook
import pytest

from solver_optimizador import (
    BiobjectiveProblem,
    ExcelExportError,
    LinearObjective,
    Sense,
    export_results_to_excel,
    solve_biobjective_epsilon_constraint,
    solve_biobjective_weighted,
    solve_multiobjective_epsilon_constraint,
)
import solver_optimizador.excel_export as excel_export_module


ROOT = Path(__file__).resolve().parents[1]
RUNNER = ROOT / "scripts" / "solve_model.py"
HYDRO = ROOT / "models" / "hidroelectrica_biobjetivo.json"
PLANNING = ROOT / "models" / "planeacion_agregada_biobjetivo.json"
MODEL_3 = ROOT / "models" / "ejemplo_tres_objetivos.json"
EXPECTED_SHEETS = [
    "Resumen",
    "Matriz_pagos",
    "Corridas",
    "Variables",
    "No_dominadas",
    "Restricciones",
]


def _runner_module():
    spec = importlib.util.spec_from_file_location("solve_model_excel", RUNNER)
    assert spec is not None and spec.loader is not None
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def _row_dicts(worksheet):
    rows = list(worksheet.iter_rows(values_only=True))
    headers = list(rows[0])
    return [dict(zip(headers, row)) for row in rows[1:]]


def _summary(worksheet):
    return {row[0]: row[1] for row in worksheet.iter_rows(min_row=2, values_only=True)}


@contextmanager
def _open_workbook(path: Path):
    workbook = load_workbook(path, data_only=True)
    try:
        yield workbook
    finally:
        workbook.close()


@pytest.fixture(scope="module")
def runner():
    return _runner_module()


@pytest.fixture(scope="module")
def hydro_solution(runner):
    name, loaded, multi = runner.load_multiobjective_model(HYDRO)
    problem = runner._as_biobjective(multi)
    result = solve_biobjective_epsilon_constraint(problem, 1, 6)
    return name, loaded, problem, result


@pytest.fixture(scope="module")
def planning_solution(runner):
    name, loaded, multi = runner.load_multiobjective_model(PLANNING)
    problem = runner._as_biobjective(multi)
    result = solve_biobjective_epsilon_constraint(problem, 1, 6)
    return name, loaded, problem, result


@pytest.fixture(scope="module")
def weighted_solution(hydro_solution):
    name, loaded, problem, _ = hydro_solution
    result = solve_biobjective_weighted(problem, num_combinations=6)
    return name, loaded, problem, result


@pytest.fixture(scope="module")
def three_objective_solution(runner):
    name, loaded, problem = runner.load_multiobjective_model(MODEL_3)
    result = solve_multiobjective_epsilon_constraint(
        problem, primary_objective=1, r_by_objective={2: 2, 3: 2}
    )
    return name, loaded, problem, result


def _export(tmp_path, bundle, method="epsilon"):
    name, _, problem, result = bundle
    output = tmp_path / f"resultado_{method}.xlsx"
    returned = export_results_to_excel(
        problem=problem,
        result=result,
        method=method,
        model_name=name,
        model_file=HYDRO,
        output_path=output,
    )
    return returned, result


def test_crea_libro_valido_con_hojas_exactas(tmp_path, hydro_solution) -> None:
    output, _ = _export(tmp_path, hydro_solution)

    assert output.is_file() and output.stat().st_size > 0
    with _open_workbook(output) as workbook:
        assert workbook.sheetnames == EXPECTED_SHEETS
        for worksheet in workbook.worksheets:
            assert worksheet.freeze_panes == "A2"
            assert worksheet.auto_filter.ref is not None
            assert all(cell.font.bold for cell in worksheet[1])


def test_resumen_contiene_metadatos_y_configuracion(tmp_path, hydro_solution) -> None:
    output, result = _export(tmp_path, hydro_solution)
    with _open_workbook(output) as workbook:
        summary = _summary(workbook["Resumen"])

    assert summary["Nombre del modelo"] == hydro_solution[0]
    assert "Generación Hidroeléctrica" in summary["Nombre del modelo"]
    assert summary["Cantidad de variables"] == 24
    assert summary["Cantidad de restricciones originales"] == 28
    assert summary["Cantidad de objetivos"] == 2
    assert summary["Objetivo principal"] == "Z1"
    assert summary["Objetivos restringidos"] == "Z2"
    assert summary["r para Z2"] == 6
    assert summary["Número total de corridas"] == len(result.runs) == 7
    assert "E_k,t" in summary["Fórmula epsilon"]
    assert summary["Operador para objetivo MAX"] == "Zk(x) >= E_k,t"


def test_matriz_de_pagos_y_corridas_coinciden_con_resultado(
    tmp_path, hydro_solution
) -> None:
    output, result = _export(tmp_path, hydro_solution)
    with _open_workbook(output) as workbook:
        payoff = _row_dicts(workbook["Matriz_pagos"])
        runs = _row_dicts(workbook["Corridas"])

    assert len(payoff) == len(result.payoff_matrix) == 2
    for row in payoff:
        expected = result.payoff_matrix[row["ancla"]]
        assert row["Z1"] == pytest.approx(expected["Z1"], abs=1e-9)
        assert row["Z2"] == pytest.approx(expected["Z2"], abs=1e-9)
    assert len(runs) == len(result.runs) == 7
    for row, expected in zip(runs, result.runs):
        assert row["t"] == expected["t"]
        assert row["E"] == pytest.approx(expected["E"], abs=1e-12)
        assert row["Z1"] == pytest.approx(expected["Z1"], abs=1e-9)
        assert row["Z2"] == pytest.approx(expected["Z2"], abs=1e-9)


def test_variables_contiene_una_columna_por_variable_y_valores_numericos(
    tmp_path, hydro_solution
) -> None:
    output, result = _export(tmp_path, hydro_solution)
    with _open_workbook(output) as workbook:
        rows = _row_dicts(workbook["Variables"])

    assert len(rows) == 7
    assert all(variable in rows[0] for variable in hydro_solution[2].variables)
    assert len(hydro_solution[2].variables) == 24
    for row, run in zip(rows, result.runs):
        for variable in hydro_solution[2].variables:
            assert isinstance(row[variable], (int, float))
            assert row[variable] == pytest.approx(run["x"][variable], abs=1e-12)


def test_precision_de_niveles_no_se_redondea(tmp_path, planning_solution) -> None:
    output, result = _export(tmp_path, planning_solution)
    with _open_workbook(output) as workbook:
        rows = _row_dicts(workbook["Corridas"])

    assert [row["E"] for row in rows] == pytest.approx(
        result.epsilon_levels, abs=1e-12
    )
    assert rows[1]["E"] == result.epsilon_levels[1]
    assert isinstance(rows[1]["E"], float)


def test_planeacion_min_min_conserva_resultados_y_catorce_variables(
    tmp_path, planning_solution
) -> None:
    output, result = _export(tmp_path, planning_solution)
    with _open_workbook(output) as workbook:
        summary = _summary(workbook["Resumen"])
        rows = _row_dicts(workbook["Variables"])

    assert len(planning_solution[2].variables) == 14
    assert summary["Objetivo Z1"].endswith("(MIN)")
    assert summary["Objetivo Z2"].endswith("(MIN)")
    assert len(rows) == len(result.runs) == 7
    assert result.epsilon_levels == pytest.approx(
        [0, 100 / 3, 200 / 3, 100, 400 / 3, 500 / 3, 200]
    )
    assert (rows[0]["Z1"], rows[0]["Z2"]) == pytest.approx((38410, 0))
    assert (rows[-1]["Z1"], rows[-1]["Z2"]) == pytest.approx((37740, 200))


def test_no_dominadas_provienen_exactamente_del_resultado(
    tmp_path, hydro_solution
) -> None:
    output, result = _export(tmp_path, hydro_solution)
    with _open_workbook(output) as workbook:
        rows = _row_dicts(workbook["No_dominadas"])

    assert [row["id"] for row in rows] == [
        solution["id"] for solution in result.nondominated_solutions
    ]
    for row, solution in zip(rows, result.nondominated_solutions):
        assert row["Z1"] == pytest.approx(solution["Z1"], abs=1e-9)
        assert row["Z2"] == pytest.approx(solution["Z2"], abs=1e-9)


def test_restricciones_se_evalúan_desde_x_sin_resolver(tmp_path, hydro_solution) -> None:
    before = deepcopy(hydro_solution[3].runs)
    output, result = _export(tmp_path, hydro_solution)
    with _open_workbook(output) as workbook:
        rows = _row_dicts(workbook["Restricciones"])

    assert result.runs == before
    assert len(rows) == 7 * (28 + 1)
    assert {row["tipo"] for row in rows} == {"original", "epsilon"}
    assert all(isinstance(row["lhs"], (int, float)) for row in rows)
    assert all(isinstance(row["holgura"], (int, float)) for row in rows)
    assert all(isinstance(row["activa"], bool) for row in rows)


def test_corrida_infactible_conserva_fila_y_variables_vacias(
    tmp_path, hydro_solution
) -> None:
    name, loaded, problem, original = hydro_solution
    result = deepcopy(original)
    result.runs[2].update(status="infeasible", x=None, Z1=None, Z2=None)
    output, _ = _export(tmp_path, (name, loaded, problem, result))
    with _open_workbook(output) as workbook:
        run_rows = _row_dicts(workbook["Corridas"])
        variable_rows = _row_dicts(workbook["Variables"])

    assert len(run_rows) == len(variable_rows) == 7
    assert run_rows[2]["estado"] == "infeasible"
    assert all(variable_rows[2][variable] is None for variable in problem.variables)


def test_ponderaciones_incluye_pesos_normalizados_y_pareto(
    tmp_path, weighted_solution
) -> None:
    output, result = _export(tmp_path, weighted_solution, method="weighted")
    with _open_workbook(output) as workbook:
        rows = _row_dicts(workbook["Corridas"])
        nondominated = _row_dicts(workbook["No_dominadas"])
        summary = _summary(workbook["Resumen"])

    assert len(rows) == len(result.weighted_runs) == 6
    assert {"alpha1", "alpha2", "N1", "N2", "W"} <= rows[0].keys()
    for row, run in zip(rows, result.weighted_runs):
        for key in ("alpha1", "alpha2", "N1", "N2", "W"):
            assert row[key] == pytest.approx(run[key], abs=1e-12)
    expected = [s for s in result.unique_solutions if s["pareto_status"] == "No dominada"]
    assert [row["id"] for row in nondominated] == [s["id"] for s in expected]
    assert summary["Número de combinaciones"] == 6


def test_soporta_sentidos_max_min(tmp_path, hydro_solution) -> None:
    name, loaded, original, _ = hydro_solution
    problem = BiobjectiveProblem(
        variables=list(original.variables),
        objective1=LinearObjective(
            name=original.objective1.name,
            sense=Sense.MAXIMIZE,
            coefficients=dict(original.objective1.coefficients),
        ),
        objective2=LinearObjective(
            name=original.objective2.name,
            sense=Sense.MINIMIZE,
            coefficients=dict(original.objective2.coefficients),
        ),
        constraints=list(original.constraints),
    )
    result = solve_biobjective_epsilon_constraint(problem, 1, 1)
    output, _ = _export(tmp_path, (name, loaded, problem, result))
    with _open_workbook(output) as workbook:
        summary = _summary(workbook["Resumen"])

    assert summary["Objetivo Z1"].endswith("(MAX)")
    assert summary["Objetivo Z2"].endswith("(MIN)")
    assert summary["Operador para objetivo MIN"] == "Zk(x) <= E_k,t"


def test_tres_objetivos_crea_columnas_dinamicas(
    tmp_path, three_objective_solution
) -> None:
    output, result = _export(tmp_path, three_objective_solution)
    with _open_workbook(output) as workbook:
        rows = _row_dicts(workbook["Corridas"])
        nondominated = _row_dicts(workbook["No_dominadas"])

    assert len(rows) == result.total_runs == 9
    assert {"t_Z2", "E_Z2", "t_Z3", "E_Z3", "Z1", "Z2", "Z3"} <= rows[0].keys()
    assert [row["id"] for row in nondominated] == [
        solution["id"] for solution in result.nondominated_solutions
    ]


def test_no_excel_no_invoca_exportador(runner, monkeypatch, hydro_solution) -> None:
    def fail_if_called(**kwargs):
        raise AssertionError("--no-excel debe omitir la exportación")

    monkeypatch.setattr(runner, "export_results_to_excel", fail_if_called)
    code = runner.main(
        [str(HYDRO), "--method", "epsilon", "--primary", "1", "--r", "1", "--no-plot", "--no-excel"]
    )
    assert code == 0


def test_no_plot_mantiene_excel_activo(tmp_path, runner, monkeypatch) -> None:
    monkeypatch.setattr(runner, "PROJECT_ROOT", tmp_path)
    code = runner.main(
        [str(HYDRO), "--method", "epsilon", "--primary", "1", "--r", "1", "--no-plot"]
    )

    assert code == 0
    assert (tmp_path / "results" / "hidroelectrica_biobjetivo_epsilon.xlsx").is_file()
    assert not list((tmp_path / "results").glob("*.png"))


def test_error_excel_se_reporta_separado_del_solver(
    runner, monkeypatch, capsys
) -> None:
    def fail_export(**kwargs):
        raise ExcelExportError("destino sin permisos")

    monkeypatch.setattr(runner, "export_results_to_excel", fail_export)
    code = runner.main(
        [str(HYDRO), "--method", "epsilon", "--primary", "1", "--r", "1", "--no-plot"]
    )
    captured = capsys.readouterr()

    assert code == 1
    assert "ERROR al exportar el libro Excel" in captured.err
    assert "destino sin permisos" in captured.err
    assert "ERROR durante la resolución" not in captured.err


def test_error_de_escritura_api_es_especifico(
    tmp_path, hydro_solution, monkeypatch
) -> None:
    def fail_save(self, filename):
        raise PermissionError("archivo abierto")

    monkeypatch.setattr(excel_export_module.Workbook, "save", fail_save)
    with pytest.raises(ExcelExportError, match="archivo abierto"):
        _export(tmp_path, hydro_solution)


def test_help_documenta_no_excel(runner) -> None:
    help_text = runner.build_parser().format_help()
    assert "--no-excel" in help_text
    assert "libro Excel" in help_text
