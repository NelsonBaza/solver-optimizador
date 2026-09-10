"""Exportación estructurada de resultados multiobjetivo a libros Excel."""

from __future__ import annotations

from enum import Enum
from pathlib import Path
from typing import Any, Iterable, Mapping, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .epsilon_constraint import EpsilonConstraintSolution
from .lp_models import (
    BiobjectiveProblem,
    LinearConstraint,
    MultiobjectiveProblem,
    MultiobjectiveSolution,
    Operator,
)
from .multiobjective_epsilon import MultiobjectiveEpsilonSolution


SHEET_NAMES = (
    "Resumen",
    "Matriz_pagos",
    "Corridas",
    "Variables",
    "No_dominadas",
    "Restricciones",
)
NUMBER_FORMAT = "0.###############"
HEADER_FILL = PatternFill("solid", fgColor="D9EAF7")


class ExcelExportError(RuntimeError):
    """Error específico al construir o guardar el libro de resultados."""


def _objectives(problem: BiobjectiveProblem | MultiobjectiveProblem) -> list[Any]:
    if isinstance(problem, BiobjectiveProblem):
        return [problem.objective1, problem.objective2]
    return list(problem.objectives)


def _runs(
    result: EpsilonConstraintSolution
    | MultiobjectiveEpsilonSolution
    | MultiobjectiveSolution,
    method: str,
) -> list[dict[str, Any]]:
    return list(result.weighted_runs if method == "weighted" else result.runs)


def _status_text(value: Any) -> str:
    if isinstance(value, Enum):
        return str(value.value)
    return "" if value is None else str(value)


def _is_optimal(status: Any) -> bool:
    normalized = _status_text(status).strip().lower()
    return normalized in {"optimal", "optimo", "óptimo"}


def _objective_values(item: Mapping[str, Any], labels: Sequence[str]) -> list[Any]:
    values = item.get("objective_values")
    if isinstance(values, Mapping):
        return [values.get(label) for label in labels]
    return [item.get(label) for label in labels]


def _display_run_number(run: Mapping[str, Any], method: str) -> int:
    raw = int(run.get("run_index", 0))
    return raw if method == "weighted" else raw + 1


def _append_table(
    worksheet: Worksheet,
    headers: Sequence[str],
    rows: Iterable[Sequence[Any]],
) -> None:
    worksheet.append(list(headers))
    for row in rows:
        worksheet.append(list(row))


def _format_sheet(worksheet: Worksheet) -> None:
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    for cell in worksheet[1]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="top", wrap_text=True)

    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if isinstance(cell.value, (int, float)) and not isinstance(
                cell.value, bool
            ):
                cell.number_format = NUMBER_FORMAT

    for column_index, cells in enumerate(worksheet.columns, start=1):
        width = max(
            10,
            min(
                50,
                max(
                    len(str(cell.value)) if cell.value is not None else 0
                    for cell in cells
                )
                + 2,
            ),
        )
        worksheet.column_dimensions[get_column_letter(column_index)].width = width


def _summary_rows(
    *,
    problem: BiobjectiveProblem | MultiobjectiveProblem,
    result: EpsilonConstraintSolution
    | MultiobjectiveEpsilonSolution
    | MultiobjectiveSolution,
    method: str,
    model_name: str,
    model_file: Path,
    tolerance: float,
) -> list[tuple[str, Any]]:
    objectives = _objectives(problem)
    runs = _runs(result, method)
    optimal_count = sum(_is_optimal(run.get("status")) for run in runs)
    nondominated = getattr(result, "nondominated_solutions", None)
    if nondominated is None:
        nondominated = [
            solution
            for solution in result.unique_solutions
            if solution.get("pareto_status") == "No dominada"
        ]
    rows: list[tuple[str, Any]] = [
        ("Nombre del modelo", model_name),
        ("Archivo fuente", str(model_file)),
        (
            "Método",
            "Método de las restricciones (epsilon-constraint)"
            if method == "epsilon"
            else "Método de ponderaciones normalizadas",
        ),
        ("Cantidad de variables", len(problem.variables)),
        ("Cantidad de restricciones originales", len(problem.constraints)),
        ("Cantidad de objetivos", len(objectives)),
    ]
    for index, objective in enumerate(objectives, start=1):
        rows.append(
            (
                f"Objetivo Z{index}",
                f"{objective.name} ({objective.sense.value.upper()})",
            )
        )

    if method == "epsilon":
        primary = result.primary_objective
        constrained = (
            [result.constrained_objective]
            if isinstance(result, EpsilonConstraintSolution)
            else list(result.constrained_objectives)
        )
        rows.extend(
            [
                ("Objetivo principal", f"Z{primary}"),
                (
                    "Objetivos restringidos",
                    ", ".join(f"Z{index}" for index in constrained),
                ),
            ]
        )
        if isinstance(result, EpsilonConstraintSolution):
            rows.extend(
                [
                    (f"r para Z{result.constrained_objective}", result.r),
                    (
                        f"Niveles para Z{result.constrained_objective}",
                        len(result.epsilon_levels),
                    ),
                ]
            )
        else:
            for index in constrained:
                rows.extend(
                    [
                        (f"r para Z{index}", result.r_by_objective[index]),
                        (
                            f"Niveles para Z{index}",
                            len(result.epsilon_levels[f"Z{index}"]),
                        ),
                    ]
                )
        rows.extend(
            [
                (
                    "Fórmula epsilon",
                    "E_k,t = Zk_min + (t/r_k)(Zk_max - Zk_min)",
                ),
                ("Operador para objetivo MAX", "Zk(x) >= E_k,t"),
                ("Operador para objetivo MIN", "Zk(x) <= E_k,t"),
            ]
        )
    else:
        rows.extend(
            [
                ("Número de combinaciones", len(runs)),
                (
                    "Pesos",
                    "alpha1 y alpha2 son no negativos y suman 1; "
                    "W = alpha1*N1 + alpha2*N2",
                ),
            ]
        )
        for key, value in result.normalization_ranges.items():
            rows.append((f"Rango de normalización {key}", value))

    rows.extend(
        [
            ("Número total de corridas", len(runs)),
            ("Corridas óptimas", optimal_count),
            ("Corridas infactibles/no óptimas", len(runs) - optimal_count),
            ("Soluciones únicas", len(result.unique_solutions)),
            ("Soluciones no dominadas obtenidas", len(nondominated)),
            (
                "Conjunto Pareto reportado",
                "Soluciones no dominadas obtenidas por el barrido",
            ),
            ("Tolerancia de evaluación", tolerance),
        ]
    )
    for key, value in result.timing.items():
        rows.append((f"Tiempo {key}", value))
    return rows


def _payoff_rows(
    problem: BiobjectiveProblem | MultiobjectiveProblem,
    result: EpsilonConstraintSolution
    | MultiobjectiveEpsilonSolution
    | MultiobjectiveSolution,
) -> tuple[list[str], list[list[Any]]]:
    labels = [f"Z{index}" for index in range(1, len(_objectives(problem)) + 1)]
    headers = ["ancla", "estado", *labels, *problem.variables]
    rows: list[list[Any]] = []
    for anchor_name, anchor in result.payoff_matrix.items():
        status = anchor.get("status")
        if status is None:
            individual = result.individual_optima.get(
                anchor_name.removeprefix("opt_") + "_opt"
            )
            status = individual.get("status") if individual else "optimal"
        x_values = anchor.get("x") or {}
        rows.append(
            [
                anchor_name,
                _status_text(status),
                *_objective_values(anchor, labels),
                *(x_values.get(variable) for variable in problem.variables),
            ]
        )
    return headers, rows


def _epsilon_columns(
    result: EpsilonConstraintSolution | MultiobjectiveEpsilonSolution,
) -> list[str]:
    if isinstance(result, EpsilonConstraintSolution):
        return ["t", "E"]
    columns: list[str] = []
    for index in result.constrained_objectives:
        columns.extend([f"t_Z{index}", f"E_Z{index}"])
    return columns


def _epsilon_values(
    run: Mapping[str, Any],
    result: EpsilonConstraintSolution | MultiobjectiveEpsilonSolution,
) -> list[Any]:
    if isinstance(result, EpsilonConstraintSolution):
        return [run.get("t"), run.get("E")]
    values: list[Any] = []
    for index in result.constrained_objectives:
        label = f"Z{index}"
        values.extend(
            [run.get("epsilon_indices", {}).get(label), run.get("epsilon_levels", {}).get(label)]
        )
    return values


def _run_rows(
    problem: BiobjectiveProblem | MultiobjectiveProblem,
    result: EpsilonConstraintSolution
    | MultiobjectiveEpsilonSolution
    | MultiobjectiveSolution,
    method: str,
) -> tuple[list[str], list[list[Any]]]:
    labels = [f"Z{index}" for index in range(1, len(_objectives(problem)) + 1)]
    runs = _runs(result, method)
    if method == "weighted":
        headers = [
            "corrida",
            "alpha1",
            "alpha2",
            "estado",
            *labels,
            "N1",
            "N2",
            "W",
            "tiempo_s",
        ]
        rows = [
            [
                _display_run_number(run, method),
                run.get("alpha1"),
                run.get("alpha2"),
                _status_text(run.get("status")),
                *_objective_values(run, labels),
                run.get("N1"),
                run.get("N2"),
                run.get("W"),
                run.get("execution_time_sec"),
            ]
            for run in runs
        ]
        return headers, rows

    epsilon_result = result
    headers = ["corrida", *_epsilon_columns(epsilon_result)]
    if isinstance(epsilon_result, EpsilonConstraintSolution):
        headers.extend(
            ["objetivo_principal", "objetivo_restringido", "operador_epsilon"]
        )
    headers.extend(["estado", *labels, "tiempo_s"])
    rows: list[list[Any]] = []
    for run in runs:
        row = [
            _display_run_number(run, method),
            *_epsilon_values(run, epsilon_result),
        ]
        if isinstance(epsilon_result, EpsilonConstraintSolution):
            row.extend(
                [
                    f"Z{run.get('primary_objective')}",
                    f"Z{run.get('constrained_objective')}",
                    run.get("constraint_operator"),
                ]
            )
        row.extend(
            [
                _status_text(run.get("status")),
                *_objective_values(run, labels),
                run.get("execution_time_sec"),
            ]
        )
        rows.append(row)
    return headers, rows


def _variable_rows(
    problem: BiobjectiveProblem | MultiobjectiveProblem,
    result: EpsilonConstraintSolution
    | MultiobjectiveEpsilonSolution
    | MultiobjectiveSolution,
    method: str,
) -> tuple[list[str], list[list[Any]]]:
    labels = [f"Z{index}" for index in range(1, len(_objectives(problem)) + 1)]
    level_headers = (
        [] if method == "weighted" else _epsilon_columns(result)
    )
    headers = ["corrida", "estado", *level_headers, *labels, *problem.variables]
    rows: list[list[Any]] = []
    for run in _runs(result, method):
        x_values = run.get("x") or {}
        levels = [] if method == "weighted" else _epsilon_values(run, result)
        rows.append(
            [
                _display_run_number(run, method),
                _status_text(run.get("status")),
                *levels,
                *_objective_values(run, labels),
                *(x_values.get(variable) for variable in problem.variables),
            ]
        )
    return headers, rows


def _join_run_indices(indices: Iterable[int], method: str) -> str:
    return ", ".join(
        str(index if method == "weighted" else index + 1) for index in indices
    )


def _generator_description(solution: Mapping[str, Any], method: str) -> str:
    if method == "weighted":
        return "; ".join(
            f"alpha1={item['alpha1']}, alpha2={item['alpha2']}"
            for item in solution.get("generated_by_weights", [])
        )
    if "epsilon_levels" in solution:
        descriptions: list[str] = []
        for level in solution["epsilon_levels"]:
            if isinstance(level, Mapping):
                descriptions.append(
                    ", ".join(f"{label}={value}" for label, value in level.items())
                )
            else:
                descriptions.append(f"E={level}")
        return "; ".join(descriptions)
    return ""


def _nondominated_rows(
    problem: BiobjectiveProblem | MultiobjectiveProblem,
    result: EpsilonConstraintSolution
    | MultiobjectiveEpsilonSolution
    | MultiobjectiveSolution,
    method: str,
) -> tuple[list[str], list[list[Any]]]:
    labels = [f"Z{index}" for index in range(1, len(_objectives(problem)) + 1)]
    selected = getattr(result, "nondominated_solutions", None)
    if selected is None:
        selected = [
            solution
            for solution in result.unique_solutions
            if solution.get("pareto_status") == "No dominada"
        ]
    headers = [
        "id",
        "corridas_generadoras",
        "repeticiones",
        "niveles_o_pesos_generadores",
        *labels,
        *problem.variables,
    ]
    rows: list[list[Any]] = []
    for solution in selected:
        run_indices = solution.get("run_indices", [])
        if method == "weighted" and not run_indices:
            weights = {
                (item["alpha1"], item["alpha2"])
                for item in solution.get("generated_by_weights", [])
            }
            run_indices = [
                run["run_index"]
                for run in result.weighted_runs
                if (run.get("alpha1"), run.get("alpha2")) in weights
            ]
        x_values = solution.get("x") or {}
        rows.append(
            [
                solution.get("id"),
                _join_run_indices(run_indices, method),
                solution.get("count", len(run_indices)),
                _generator_description(solution, method),
                *_objective_values(solution, labels),
                *(x_values.get(variable) for variable in problem.variables),
            ]
        )
    return headers, rows


def _constraint_row(
    *,
    run_number: int,
    status: str,
    constraint: LinearConstraint,
    x_values: Mapping[str, float],
    tolerance: float,
    kind: str,
) -> list[Any]:
    vector = dict(x_values)
    lhs = constraint.evaluate_lhs(vector)
    slack = constraint.calculate_slack(vector)
    return [
        run_number,
        status,
        constraint.name,
        kind,
        lhs,
        constraint.operator.value,
        constraint.rhs,
        slack,
        abs(slack) < tolerance,
    ]


def _epsilon_constraints_for_run(
    problem: BiobjectiveProblem | MultiobjectiveProblem,
    result: EpsilonConstraintSolution | MultiobjectiveEpsilonSolution,
    run: Mapping[str, Any],
) -> list[LinearConstraint]:
    objectives = _objectives(problem)
    if isinstance(result, EpsilonConstraintSolution):
        index = result.constrained_objective
        return [
            LinearConstraint(
                name=f"_epsilon_Z{index}_t_{run['t']}",
                coefficients=dict(objectives[index - 1].coefficients),
                operator=Operator.from_str(run["constraint_operator"]),
                rhs=run["E"],
            )
        ]
    constraints: list[LinearConstraint] = []
    for index in result.constrained_objectives:
        label = f"Z{index}"
        constraints.append(
            LinearConstraint(
                name=f"_epsilon_{label}_t_{run['epsilon_indices'][label]}",
                coefficients=dict(objectives[index - 1].coefficients),
                operator=Operator.from_str(run["constraint_operators"][label]),
                rhs=run["epsilon_levels"][label],
            )
        )
    return constraints


def _restriction_rows(
    problem: BiobjectiveProblem | MultiobjectiveProblem,
    result: EpsilonConstraintSolution
    | MultiobjectiveEpsilonSolution
    | MultiobjectiveSolution,
    method: str,
    tolerance: float,
) -> tuple[list[str], list[list[Any]]]:
    headers = [
        "corrida",
        "estado",
        "nombre_restriccion",
        "tipo",
        "lhs",
        "operador",
        "rhs",
        "holgura",
        "activa",
    ]
    rows: list[list[Any]] = []
    for run in _runs(result, method):
        if run.get("x") is None:
            continue
        run_number = _display_run_number(run, method)
        status = _status_text(run.get("status"))
        for constraint in problem.constraints:
            rows.append(
                _constraint_row(
                    run_number=run_number,
                    status=status,
                    constraint=constraint,
                    x_values=run["x"],
                    tolerance=tolerance,
                    kind="original",
                )
            )
        if method == "epsilon":
            for constraint in _epsilon_constraints_for_run(problem, result, run):
                rows.append(
                    _constraint_row(
                        run_number=run_number,
                        status=status,
                        constraint=constraint,
                        x_values=run["x"],
                        tolerance=tolerance,
                        kind="epsilon",
                    )
                )
    return headers, rows


def export_results_to_excel(
    *,
    problem: BiobjectiveProblem | MultiobjectiveProblem,
    result: EpsilonConstraintSolution
    | MultiobjectiveEpsilonSolution
    | MultiobjectiveSolution,
    method: str,
    model_name: str,
    model_file: str | Path,
    output_path: str | Path,
    tolerance: float = 1e-6,
) -> Path:
    """Exporta resultados existentes sin ejecutar ni modificar el solver."""

    if method not in {"epsilon", "weighted"}:
        raise ValueError("method debe ser 'epsilon' o 'weighted'.")
    if method == "weighted" and not isinstance(result, MultiobjectiveSolution):
        raise TypeError("El resultado weighted debe ser MultiobjectiveSolution.")
    if method == "epsilon" and not isinstance(
        result, (EpsilonConstraintSolution, MultiobjectiveEpsilonSolution)
    ):
        raise TypeError("El resultado epsilon no tiene el tipo esperado.")
    if tolerance <= 0:
        raise ValueError("tolerance debe ser mayor que cero.")

    source = Path(model_file)
    destination = Path(output_path)
    if destination.suffix.lower() != ".xlsx":
        raise ValueError("output_path debe tener extensión .xlsx.")
    workbook = Workbook()
    summary = workbook.active
    summary.title = SHEET_NAMES[0]
    sheets = {name: workbook.create_sheet(name) for name in SHEET_NAMES[1:]}

    _append_table(
        summary,
        ["Campo", "Valor"],
        _summary_rows(
            problem=problem,
            result=result,
            method=method,
            model_name=model_name,
            model_file=source,
            tolerance=tolerance,
        ),
    )
    headers, rows = _payoff_rows(problem, result)
    _append_table(sheets["Matriz_pagos"], headers, rows)
    headers, rows = _run_rows(problem, result, method)
    _append_table(sheets["Corridas"], headers, rows)
    headers, rows = _variable_rows(problem, result, method)
    _append_table(sheets["Variables"], headers, rows)
    headers, rows = _nondominated_rows(problem, result, method)
    _append_table(sheets["No_dominadas"], headers, rows)
    headers, rows = _restriction_rows(problem, result, method, tolerance)
    _append_table(sheets["Restricciones"], headers, rows)

    for worksheet in workbook.worksheets:
        _format_sheet(worksheet)

    try:
        destination.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(destination)
    except OSError as exc:
        raise ExcelExportError(
            f"No se pudo guardar el libro Excel en '{destination}': {exc}"
        ) from exc
    finally:
        workbook.close()
    return destination.resolve()


__all__ = ["ExcelExportError", "SHEET_NAMES", "export_results_to_excel"]
